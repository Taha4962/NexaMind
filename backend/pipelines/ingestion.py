"""
NexaMind Backend — Document Ingestion Pipeline

Handles the full ingestion process for a document:
Download from Cloudinary -> Text Extraction -> Chunking ->
Embedding -> Storing in ChromaDB -> Updating MongoDB Status.
"""

import logging
import io
import re
import httpx
import fitz  # PyMuPDF
import docx
from bson import ObjectId

from utils.chunker import chunk_text
from utils.embedder import GeminiEmbedder
from db.mongo import get_db

try:
    import chromadb
    from chromadb.config import Settings
    CHROMA_AVAILABLE = True
except ImportError:
    CHROMA_AVAILABLE = False

logger = logging.getLogger("nexamind.pipelines.ingestion")

class DocumentIngestionPipeline:
    def __init__(self) -> None:
        self.embedder = GeminiEmbedder()
        if CHROMA_AVAILABLE:
            # We use a persistent client in the same directory as the backend
            self.chroma_client = chromadb.PersistentClient(path="./chroma_db")
        else:
            logger.warning("ChromaDB not available. Vector storage will be skipped.")

    async def ingest(
        self, document_id: str, cloudinary_url: str, file_type: str, user_id: str
    ) -> dict[str, bool | int]:
        """Full pipeline: download -> extract -> chunk -> embed -> store"""
        try:
            logger.info("Starting ingestion for doc %s", document_id)
            
            # Step 1: Download file
            raw_bytes = await self._download_file(cloudinary_url)
            
            # Step 2: Extract text
            text = await self._extract_text(raw_bytes, file_type)
            
            # Step 3: Chunk text
            chunks = chunk_text(text, chunk_size=500, overlap=50)
            if not chunks:
                raise ValueError("No text could be extracted from the document.")

            # Step 4: Embed chunks
            embeddings = await self._embed_chunks(chunks)

            # Step 5: Store in Vector DB
            await self._store_in_vector_db(chunks, embeddings, document_id, user_id)

            # Step 6: Update document status
            await self._update_status(document_id, "ready", chunk_count=len(chunks))
            
            logger.info("Ingestion complete for doc %s", document_id)
            return {"success": True, "chunks": len(chunks)}
        except Exception as e:
            logger.error("Ingestion failed for doc %s: %s", document_id, str(e), exc_info=True)
            await self._update_status(document_id, "failed", error=str(e))
            raise

    async def _download_file(self, url: str) -> bytes:
        """Download file from Cloudinary URL"""
        async with httpx.AsyncClient() as client:
            response = await client.get(url, timeout=30.0)
            response.raise_for_status()
            return response.content

    async def _extract_text(self, raw_bytes: bytes, file_type: str) -> str:
        """Extract text based on file type"""
        if file_type == "pdf":
            text_parts = []
            # PyMuPDF opens from stream
            with fitz.open(stream=raw_bytes, filetype="pdf") as doc:
                for page in doc:
                    text_parts.append(page.get_text())
            return "\n".join(text_parts)
            
        elif file_type == "docx":
            file_stream = io.BytesIO(raw_bytes)
            doc = docx.Document(file_stream)
            text_parts = []
            for para in doc.paragraphs:
                text_parts.append(para.text)
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        text_parts.append(cell.text)
            return "\n".join(text_parts)
            
        elif file_type == "txt":
            try:
                return raw_bytes.decode("utf-8")
            except UnicodeDecodeError:
                return raw_bytes.decode("latin-1")
                
        elif file_type == "image":
            raise NotImplementedError("Image OCR not supported in free tier")
            
        else:
            raise ValueError(f"Unsupported file type: {file_type}")

    async def _embed_chunks(self, chunks: list[str]) -> list[list[float]]:
        """Embed all chunks using GeminiEmbedder"""
        return await self.embedder.embed_batch(chunks)

    async def _store_in_vector_db(
        self, chunks: list[str], embeddings: list[list[float]], doc_id: str, user_id: str
    ) -> None:
        """Store chunks and embeddings in ChromaDB per-user collection"""
        if not CHROMA_AVAILABLE:
            logger.warning("Skipping vector store for doc %s (ChromaDB not available)", doc_id)
            return

        # Sanitize user_id for collection name (alphanumeric and underscores)
        safe_user_id = re.sub(r'[^a-zA-Z0-9_]', '', user_id)
        collection_name = f"user_{safe_user_id}_docs"
        
        # We need to run sync chromadb calls in a thread pool or just synchronously
        # since we are inside an async method
        collection = self.chroma_client.get_or_create_collection(name=collection_name)
        
        ids = [f"{doc_id}_chunk_{i}" for i in range(len(chunks))]
        metadatas = [
            {
                "documentId": doc_id,
                "userId": user_id,
                "chunkIndex": i,
                "textPreview": chunk[:100]
            }
            for i, chunk in enumerate(chunks)
        ]
        
        collection.add(
            ids=ids,
            embeddings=embeddings,
            metadatas=metadatas,
            documents=chunks
        )

    async def _update_status(
        self, doc_id: str, status: str, chunk_count: int | None = None, error: str | None = None
    ) -> None:
        """Update MongoDB document status"""
        db = get_db()
        update_fields: dict[str, str | int] = {"status": status}
        if chunk_count is not None:
            update_fields["chunkCount"] = chunk_count
        if error is not None:
            update_fields["errorMessage"] = error
            
        await db.documents.update_one(
            {"_id": ObjectId(doc_id)},
            {"$set": update_fields}
        )
